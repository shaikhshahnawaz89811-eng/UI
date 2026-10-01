#!/usr/bin/env python3
import csv, sys
from pathlib import Path

def fail(msg):
    print('FAIL:', msg)
    return 1

if len(sys.argv) != 2:
    raise SystemExit(fail('manifest path required'))
manifest = Path(sys.argv[1])
if not manifest.is_file():
    raise SystemExit(fail('manifest missing'))
counts = {}
checked = 0
try:
    from docx import Document
    import openpyxl
    from pptx import Presentation
except Exception as e:
    raise SystemExit(fail('required office Python packages unavailable: ' + str(e)))
with manifest.open(encoding='utf-8', newline='') as f:
    next(f, None)
    for line in f:
        parts = line.rstrip('\n').split('\t')
        if len(parts) != 5:
            continue
        skill, raw = parts[3], Path(parts[4])
        path = raw
        if not path.is_file() or path.stat().st_size == 0:
            raise SystemExit(fail(f'{skill} output missing/empty: {path}'))
        if skill == 'DOCX':
            d = Document(path)
            _ = [p.text for p in d.paragraphs]
        elif skill == 'XLSX':
            w = openpyxl.load_workbook(path, read_only=True, data_only=False)
            _ = w.sheetnames
            w.close()
        elif skill == 'PPTX':
            p = Presentation(path)
            _ = len(p.slides)
        elif skill == 'PDF_CREATOR':
            if path.read_bytes()[:5] != b'%PDF-':
                raise SystemExit(fail(f'PDF header invalid: {path}'))
        counts[skill] = counts.get(skill, 0) + 1
        checked += 1
print(f'checked={checked}; ' + ', '.join(f'{k}={v}' for k,v in sorted(counts.items())))
print('PASS')
